import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import os

# ─── Couleurs SUNU ────────────────────────────────────────
BLEU = "#1A2B4A"
OR = "#C9A84C"

def charger_fichier():
    """Ouvre une boîte de dialogue pour choisir le fichier Excel."""
    chemin = filedialog.askopenfilename(
        title="Choisir le fichier de collecte",
        filetypes=[("Fichiers Excel", "*.xlsx")]
    )
    if chemin:
        entry_fichier.delete(0, tk.END)
        entry_fichier.insert(0, chemin)

def calculer():
    """Calcule les indicateurs NPS, CSAT, CES."""
    chemin = entry_fichier.get().strip()
    if not chemin or not os.path.exists(chemin):
        messagebox.showwarning("Attention", "Veuillez sélectionner un fichier Excel valide.")
        return

    # Chargement des données
    df = pd.read_excel(chemin, sheet_name="Collecte")

    # Vérification colonnes
    colonnes_requises = ["Catégorie NPS", "CSAT", "CES", "Périmètre", "Canal"]
    for col in colonnes_requises:
        if col not in df.columns:
            messagebox.showerror("Erreur", f"Colonne manquante : {col}")
            return

    # Nettoyage résultats
    for widget in frame_resultats.winfo_children():
        widget.destroy()

    # ─── Calcul global ────────────────────────────────────
    resultats = []
    for perimetre in ["VIE", "IARD", "GLOBAL"]:
        if perimetre == "GLOBAL":
            data = df
        else:
            data = df[df["Périmètre"] == perimetre]

        total = len(data)
        if total == 0:
            continue

        # NPS
        promoteurs = len(data[data["Catégorie NPS"] == "Promoteur"])
        detracteurs = len(data[data["Catégorie NPS"] == "Détracteur"])
        nps = round((promoteurs / total * 100) - (detracteurs / total * 100), 1)

        # CSAT
        satisfaits = len(data[data["CSAT"] == "Oui"])
        csat = round(satisfaits / total * 100, 1)

        # CES
        effort_eleve = len(data[data["CES"] == "Élevé"])
        effort_faible = len(data[data["CES"] == "Faible"])
        ces = round((effort_eleve / total * 100) - (effort_faible / total * 100), 1)

        resultats.append({
            "Périmètre": perimetre,
            "Total": total,
            "NPS (%)": nps,
            "CSAT (%)": csat,
            "CES (%)": ces
        })

    # ─── Affichage résultats ──────────────────────────────
    tk.Label(frame_resultats, text="Résultats des indicateurs",
             font=("Helvetica", 12, "bold"),
             bg="white", fg=BLEU).grid(row=0, column=0, columnspan=5, pady=10)

    # En-têtes
    entetes = ["Périmètre", "Total clients", "NPS (%)", "CSAT (%)", "CES (%)"]
    for col, entete in enumerate(entetes):
        tk.Label(frame_resultats, text=entete,
                 font=("Helvetica", 10, "bold"),
                 bg=BLEU, fg="white",
                 width=14, relief="ridge").grid(row=1, column=col, padx=2, pady=2)

    # Données
    for row, res in enumerate(resultats, start=2):
        valeurs = [res["Périmètre"], res["Total"],
                   res["NPS (%)"], res["CSAT (%)"], res["CES (%)"]]
        bg = "#F0F0F0" if row % 2 == 0 else "white"
        for col, val in enumerate(valeurs):
            tk.Label(frame_resultats, text=val,
                     bg=bg, width=14, relief="ridge",
                     font=("Helvetica", 10)).grid(row=row, column=col, padx=2, pady=2)

    # Seuil NPS
    tk.Label(frame_resultats,
             text="⚠️ Seuil NPS requis : 50% — En dessous = insuffisant",
             font=("Helvetica", 9, "italic"),
             bg="white", fg="red").grid(row=row+1, column=0, columnspan=5, pady=8)

    # Bouton export
    tk.Button(frame_resultats, text="📥 Exporter les résultats",
              command=lambda: exporter(resultats),
              bg=OR, fg="white",
              font=("Helvetica", 11, "bold"),
              width=25).grid(row=row+2, column=0, columnspan=5, pady=10)

def exporter(resultats):
    """Exporte les résultats dans un fichier Excel."""
    chemin = filedialog.asksaveasfilename(
        title="Enregistrer les résultats",
        defaultextension=".xlsx",
        filetypes=[("Fichiers Excel", "*.xlsx")],
        initialfile="SUNU_Indicateurs.xlsx"
    )
    if not chemin:
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Indicateurs"

    # En-têtes
    entetes = ["Périmètre", "Total clients", "NPS (%)", "CSAT (%)", "CES (%)"]
    for col, entete in enumerate(entetes, start=1):
        cell = ws.cell(row=1, column=col, value=entete)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1A2B4A")
        cell.alignment = Alignment(horizontal="center")

    # Données
    for row, res in enumerate(resultats, start=2):
        ws.cell(row=row, column=1, value=res["Périmètre"])
        ws.cell(row=row, column=2, value=res["Total"])
        ws.cell(row=row, column=3, value=res["NPS (%)"])
        ws.cell(row=row, column=4, value=res["CSAT (%)"])
        ws.cell(row=row, column=5, value=res["CES (%)"])

    # Largeur colonnes
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18

    wb.save(chemin)
    messagebox.showinfo("Succès", f"Résultats exportés avec succès ✅\n{chemin}")

# ─── Interface graphique ──────────────────────────────────
root = tk.Tk()
root.title("SUNU Assurances — Calcul des Indicateurs")
root.geometry("780x600")
root.resizable(False, False)
root.configure(bg=BLEU)

# Titre
tk.Label(root, text="SUNU Business — Calcul des Indicateurs",
         font=("Helvetica", 14, "bold"),
         bg=BLEU, fg="white").pack(pady=15)

# Sélection fichier
frame_fichier = tk.Frame(root, bg="white", padx=15, pady=15)
frame_fichier.pack(padx=20, pady=5, fill="x")

tk.Label(frame_fichier, text="Fichier de collecte :",
         bg="white", font=("Helvetica", 10, "bold")).grid(row=0, column=0, sticky="w")
entry_fichier = tk.Entry(frame_fichier, width=50)
entry_fichier.grid(row=0, column=1, padx=10)
tk.Button(frame_fichier, text="📂 Parcourir",
          command=charger_fichier,
          bg=BLEU, fg="white",
          font=("Helvetica", 10)).grid(row=0, column=2)

# Bouton calculer
tk.Button(root, text="⚙️ Calculer les indicateurs",
          command=calculer,
          bg=OR, fg="white",
          font=("Helvetica", 12, "bold"),
          width=25).pack(pady=10)

# Résultats
frame_resultats = tk.Frame(root, bg="white", padx=15, pady=15)
frame_resultats.pack(padx=20, pady=5, fill="both", expand=True)

root.mainloop()