import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import pandas as pd
import matplotlib.pyplot as plt
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
import os
from datetime import datetime

# ─── Couleurs SUNU ────────────────────────────────────────
BLEU = "#1A2B4A"
OR = "#C9A84C"
BLANC = "white"
FICHIER_EXCEL = r"C:\Users\USER\OneDrive - ENSEA\Bureau\SUNU\SUNU_Collecte.xlsx"

def initialiser_excel():
    if not os.path.exists(FICHIER_EXCEL):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Collecte"
        ws.append(["Date", "Nom Employé", "Canal", "Périmètre",
                   "NPS Score", "Catégorie NPS", "CSAT", "CES"])
        wb.save(FICHIER_EXCEL)

def classer_nps(score):
    if score >= 9:
        return "Promoteur"
    elif score >= 7:
        return "Passif"
    else:
        return "Détracteur"

def enregistrer():
    employe = entry_employe.get().strip()
    canal = combo_canal.get()
    perimetre = combo_perimetre.get()
    nps_val = entry_nps.get().strip()
    csat = csat_var.get()
    ces = ces_var.get()

    if not employe:
        messagebox.showwarning("Attention", "Veuillez entrer le nom de l'employé.")
        return
    if not canal:
        messagebox.showwarning("Attention", "Veuillez sélectionner un canal.")
        return
    if not perimetre:
        messagebox.showwarning("Attention", "Veuillez sélectionner un périmètre.")
        return
    try:
        nps_score = int(nps_val)
        if nps_score < 0 or nps_score > 10:
            raise ValueError
    except ValueError:
        messagebox.showwarning("Attention", "Le score NPS doit être entre 0 et 10.")
        return
    if not csat:
        messagebox.showwarning("Attention", "Veuillez répondre à la question CSAT.")
        return
    if not ces:
        messagebox.showwarning("Attention", "Veuillez répondre à la question CES.")
        return

    wb = openpyxl.load_workbook(FICHIER_EXCEL)
    ws = wb["Collecte"]
    ws.append([
        datetime.now().strftime("%d/%m/%Y %H:%M"),
        employe, canal, perimetre,
        nps_score, classer_nps(nps_score),
        csat, ces
    ])
    wb.save(FICHIER_EXCEL)
    messagebox.showinfo("Succès", "Réponse enregistrée avec succès ✅")
    entry_nps.delete(0, tk.END)
    csat_var.set("")
    ces_var.set("")

def calculer():
    if not os.path.exists(FICHIER_EXCEL):
        messagebox.showwarning("Attention", "Aucune donnée disponible.")
        return

    df = pd.read_excel(FICHIER_EXCEL, sheet_name="Collecte")

    for widget in frame_resultats.winfo_children():
        widget.destroy()

    resultats = []
    for perimetre in ["VIE", "IARD", "GLOBAL"]:
        data = df if perimetre == "GLOBAL" else df[df["Périmètre"] == perimetre]
        total = len(data)
        if total == 0:
            continue
        promoteurs = len(data[data["Catégorie NPS"] == "Promoteur"])
        detracteurs = len(data[data["Catégorie NPS"] == "Détracteur"])
        nps = round((promoteurs / total * 100) - (detracteurs / total * 100), 1)
        satisfaits = len(data[data["CSAT"] == "Oui"])
        csat = round(satisfaits / total * 100, 1)
        effort_eleve = len(data[data["CES"] == "Élevé"])
        effort_faible = len(data[data["CES"] == "Faible"])
        ces = round((effort_eleve / total * 100) - (effort_faible / total * 100), 1)
        resultats.append({"Périmètre": perimetre, "Total": total,
                          "NPS (%)": nps, "CSAT (%)": csat, "CES (%)": ces})

    tk.Label(frame_resultats, text="Résultats des indicateurs",
             font=("Helvetica", 12, "bold"), bg=BLANC, fg=BLEU).grid(
             row=0, column=0, columnspan=5, pady=10)

    entetes = ["Périmètre", "Total clients", "NPS (%)", "CSAT (%)", "CES (%)"]
    for col, entete in enumerate(entetes):
        tk.Label(frame_resultats, text=entete,
                 font=("Helvetica", 10, "bold"),
                 bg=BLEU, fg=BLANC, width=14, relief="ridge").grid(
                 row=1, column=col, padx=2, pady=2)

    for row, res in enumerate(resultats, start=2):
        valeurs = [res["Périmètre"], res["Total"],
                   res["NPS (%)"], res["CSAT (%)"], res["CES (%)"]]
        bg = "#F0F0F0" if row % 2 == 0 else BLANC
        for col, val in enumerate(valeurs):
            tk.Label(frame_resultats, text=val, bg=bg,
                     width=14, relief="ridge",
                     font=("Helvetica", 10)).grid(row=row, column=col, padx=2, pady=2)

    tk.Label(frame_resultats,
             text="⚠️ Seuil NPS requis : 50% — En dessous = insuffisant",
             font=("Helvetica", 9, "italic"),
             bg=BLANC, fg="red").grid(row=row+1, column=0, columnspan=5, pady=8)

    tk.Button(frame_resultats, text="📥 Exporter les résultats",
              command=lambda: exporter(resultats),
              bg=OR, fg=BLANC,
              font=("Helvetica", 11, "bold"),
              width=25).grid(row=row+2, column=0, columnspan=5, pady=10)

def exporter(resultats):
    chemin = filedialog.asksaveasfilename(
        title="Enregistrer les résultats",
        defaultextension=".xlsx",
        filetypes=[("Fichiers Excel", "*.xlsx")],
        initialfile="SUNU_Indicateurs.xlsx"
    )
    if not chemin:
        return
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Indicateurs"
    entetes = ["Périmètre", "Total clients", "NPS (%)", "CSAT (%)", "CES (%)"]
    for col, entete in enumerate(entetes, start=1):
        cell = ws.cell(row=1, column=col, value=entete)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1A2B4A")
        cell.alignment = Alignment(horizontal="center")
    for row, res in enumerate(resultats, start=2):
        ws.cell(row=row, column=1, value=res["Périmètre"])
        ws.cell(row=row, column=2, value=res["Total"])
        ws.cell(row=row, column=3, value=res["NPS (%)"])
        ws.cell(row=row, column=4, value=res["CSAT (%)"])
        ws.cell(row=row, column=5, value=res["CES (%)"])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18
    wb.save(chemin)
    messagebox.showinfo("Succès", f"Résultats exportés ✅\n{chemin}")

def afficher_dashboard():
    if not os.path.exists(FICHIER_EXCEL):
        messagebox.showwarning("Attention", "Aucune donnée disponible.")
        return

    df = pd.read_excel(FICHIER_EXCEL, sheet_name="Collecte")

    for widget in frame_dashboard.winfo_children():
        widget.destroy()

    resultats = {}
    for p in ["VIE", "IARD", "GLOBAL"]:
        data = df if p == "GLOBAL" else df[df["Périmètre"] == p]
        total = len(data)
        if total == 0:
            continue
        promoteurs = len(data[data["Catégorie NPS"] == "Promoteur"])
        detracteurs = len(data[data["Catégorie NPS"] == "Détracteur"])
        nps = round((promoteurs/total*100) - (detracteurs/total*100), 1)
        satisfaits = len(data[data["CSAT"] == "Oui"])
        csat = round(satisfaits/total*100, 1)
        effort_eleve = len(data[data["CES"] == "Élevé"])
        effort_faible = len(data[data["CES"] == "Faible"])
        ces = round((effort_eleve/total*100) - (effort_faible/total*100), 1)
        resultats[p] = {"NPS": nps, "CSAT": csat, "CES": ces}

    if not resultats:
        messagebox.showwarning("Attention", "Aucune donnée disponible.")
        return

    labels = list(resultats.keys())
    nps_vals = [resultats[p]["NPS"] for p in labels]
    csat_vals = [resultats[p]["CSAT"] for p in labels]
    ces_vals = [resultats[p]["CES"] for p in labels]

    fig, axes = plt.subplots(1, 3, figsize=(11, 3.5))
    fig.patch.set_facecolor("#F5F5F5")
    couleurs = ["#2E6DA4", "#C9A84C", "#1A2B4A"]

    axes[0].bar(labels, nps_vals, color=couleurs)
    axes[0].set_title("NPS (%)", fontsize=12, fontweight="bold", color=BLEU)
    axes[0].axhline(y=50, color="red", linestyle="--", label="Seuil 50%")
    axes[0].legend(fontsize=8)
    axes[0].set_ylim(-100, 100)
    for i, val in enumerate(nps_vals):
        axes[0].text(i, val+1, f"{val}%", ha="center", fontsize=9, fontweight="bold")

    axes[1].bar(labels, csat_vals, color=couleurs)
    axes[1].set_title("CSAT (%)", fontsize=12, fontweight="bold", color=BLEU)
    axes[1].set_ylim(0, 120)
    for i, val in enumerate(csat_vals):
        axes[1].text(i, val+1, f"{val}%", ha="center", fontsize=9, fontweight="bold")

    axes[2].bar(labels, ces_vals, color=couleurs)
    axes[2].set_title("CES (%)", fontsize=12, fontweight="bold", color=BLEU)
    axes[2].axhline(y=0, color="red", linestyle="--", label="Seuil 0%")
    axes[2].legend(fontsize=8)
    axes[2].set_ylim(-100, 100)
    for i, val in enumerate(ces_vals):
        axes[2].text(i, val+1, f"{val}%", ha="center", fontsize=9, fontweight="bold")

    plt.tight_layout()

    canvas = FigureCanvasTkAgg(fig, master=frame_dashboard)
    canvas.draw()
    canvas.get_tk_widget().pack(fill="both", expand=True)

    tk.Button(frame_dashboard, text="📥 Exporter le dashboard (image)",
              command=lambda: exporter_image(fig),
              bg=OR, fg=BLANC,
              font=("Helvetica", 11, "bold"),
              width=30).pack(pady=10)

def exporter_image(fig):
    chemin = filedialog.asksaveasfilename(
        title="Enregistrer le dashboard",
        defaultextension=".png",
        filetypes=[("Image PNG", "*.png")],
        initialfile="SUNU_Dashboard.png"
    )
    if chemin:
        fig.savefig(chemin, dpi=150, bbox_inches="tight")
        messagebox.showinfo("Succès", f"Dashboard exporté ✅\n{chemin}")

# ─── Interface principale ─────────────────────────────────
initialiser_excel()

root = tk.Tk()
root.title("SUNU Assurances — Relation Client")
root.geometry("800x650")
root.resizable(False, False)
root.configure(bg=BLEU)

# Titre
tk.Label(root, text="SUNU Business — Relation Client",
         font=("Helvetica", 15, "bold"),
         bg=BLEU, fg=BLANC).pack(pady=10)

# ─── Onglets ──────────────────────────────────────────────
style = ttk.Style()
style.theme_use("default")
style.configure("TNotebook", background=BLEU, borderwidth=0)
style.configure("TNotebook.Tab", background="#2E6DA4", foreground=BLANC,
                font=("Helvetica", 11, "bold"), padding=[20, 8])
style.map("TNotebook.Tab", background=[("selected", OR)],
          foreground=[("selected", BLANC)])

notebook = ttk.Notebook(root)
notebook.pack(fill="both", expand=True, padx=15, pady=5)

# ══════════════════════════════════════════════════════════
# ONGLET 1 — SAISIE
# ══════════════════════════════════════════════════════════
tab_saisie = tk.Frame(notebook, bg=BLANC)
notebook.add(tab_saisie, text="📝  Saisie Client")

frame_saisie = tk.Frame(tab_saisie, bg=BLANC, padx=30, pady=20)
frame_saisie.pack(fill="both", expand=True)

tk.Label(frame_saisie, text="Nom de l'employé :",
         bg=BLANC, font=("Helvetica", 11, "bold")).grid(row=0, column=0, sticky="w", pady=10)
entry_employe = tk.Entry(frame_saisie, width=35, font=("Helvetica", 11))
entry_employe.grid(row=0, column=1, pady=10)

tk.Label(frame_saisie, text="Canal :",
         bg=BLANC, font=("Helvetica", 11, "bold")).grid(row=1, column=0, sticky="w", pady=10)
combo_canal = ttk.Combobox(frame_saisie, values=["Appel", "QR Code"],
                           state="readonly", width=33, font=("Helvetica", 11))
combo_canal.grid(row=1, column=1, pady=10)

tk.Label(frame_saisie, text="Périmètre :",
         bg=BLANC, font=("Helvetica", 11, "bold")).grid(row=2, column=0, sticky="w", pady=10)
combo_perimetre = ttk.Combobox(frame_saisie, values=["VIE", "IARD"],
                               state="readonly", width=33, font=("Helvetica", 11))
combo_perimetre.grid(row=2, column=1, pady=10)

tk.Label(frame_saisie, text="Score NPS (0 à 10) :",
         bg=BLANC, font=("Helvetica", 11, "bold")).grid(row=3, column=0, sticky="w", pady=10)
entry_nps = tk.Entry(frame_saisie, width=35, font=("Helvetica", 11))
entry_nps.grid(row=3, column=1, pady=10)

tk.Label(frame_saisie, text="CSAT — Satisfait ?",
         bg=BLANC, font=("Helvetica", 11, "bold")).grid(row=4, column=0, sticky="w", pady=10)
csat_var = tk.StringVar()
csat_frame = tk.Frame(frame_saisie, bg=BLANC)
csat_frame.grid(row=4, column=1, sticky="w")
tk.Radiobutton(csat_frame, text="Oui", variable=csat_var, value="Oui",
               bg=BLANC, font=("Helvetica", 11)).pack(side="left", padx=10)
tk.Radiobutton(csat_frame, text="Non", variable=csat_var, value="Non",
               bg=BLANC, font=("Helvetica", 11)).pack(side="left")

tk.Label(frame_saisie, text="CES — Effort fourni :",
         bg=BLANC, font=("Helvetica", 11, "bold")).grid(row=5, column=0, sticky="w", pady=10)
ces_var = tk.StringVar()
ces_frame = tk.Frame(frame_saisie, bg=BLANC)
ces_frame.grid(row=5, column=1, sticky="w")
tk.Radiobutton(ces_frame, text="Faible", variable=ces_var, value="Faible",
               bg=BLANC, font=("Helvetica", 11)).pack(side="left", padx=5)
tk.Radiobutton(ces_frame, text="Modéré", variable=ces_var, value="Modéré",
               bg=BLANC, font=("Helvetica", 11)).pack(side="left", padx=5)
tk.Radiobutton(ces_frame, text="Élevé", variable=ces_var, value="Élevé",
               bg=BLANC, font=("Helvetica", 11)).pack(side="left", padx=5)

tk.Button(frame_saisie, text="✅  Enregistrer",
          command=enregistrer,
          bg=OR, fg=BLANC,
          font=("Helvetica", 13, "bold"),
          width=25).grid(row=6, column=0, columnspan=2, pady=25)

# ══════════════════════════════════════════════════════════
# ONGLET 2 — INDICATEURS
# ══════════════════════════════════════════════════════════
tab_indicateurs = tk.Frame(notebook, bg=BLANC)
notebook.add(tab_indicateurs, text="⚙️  Indicateurs")

tk.Button(tab_indicateurs, text="⚙️  Calculer les indicateurs",
          command=calculer,
          bg=BLEU, fg=BLANC,
          font=("Helvetica", 12, "bold"),
          width=25).pack(pady=20)

frame_resultats = tk.Frame(tab_indicateurs, bg=BLANC)
frame_resultats.pack(fill="both", expand=True, padx=20)

# ══════════════════════════════════════════════════════════
# ONGLET 3 — DASHBOARD
# ══════════════════════════════════════════════════════════
tab_dashboard = tk.Frame(notebook, bg=BLANC)
notebook.add(tab_dashboard, text="📊  Dashboard")

tk.Button(tab_dashboard, text="📊  Afficher le Dashboard",
          command=afficher_dashboard,
          bg=BLEU, fg=BLANC,
          font=("Helvetica", 12, "bold"),
          width=25).pack(pady=20)

frame_dashboard = tk.Frame(tab_dashboard, bg=BLANC)
frame_dashboard.pack(fill="both", expand=True, padx=20)

root.mainloop()